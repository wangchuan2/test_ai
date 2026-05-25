"""Generate user manual document for Test Case Generator."""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, color):
    """Set cell background color."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_run_font(run, name='微软雅黑', size=10.5, bold=False, color=None):
    """Set run font properties."""
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_custom(doc, text, level=1):
    """Add styled heading."""
    para = doc.add_heading(text, level=level)
    for run in para.runs:
        set_run_font(run, size=(18 if level == 1 else 14 if level == 2 else 12),
                     bold=True, color=(0x44, 0x72, 0xC4) if level == 1 else (0x33, 0x33, 0x33))
    return para


def add_paragraph_custom(doc, text, bold=False, size=10.5, indent=True):
    """Add styled paragraph."""
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.first_line_indent = Cm(0.74)
    para.paragraph_format.line_spacing = 1.5
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return para


def add_code_block(doc, code):
    """Add code block styled paragraph."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(1)
    para.paragraph_format.line_spacing = 1.3
    run = para.add_run(code)
    set_run_font(run, name='Consolas', size=9, color=(0x33, 0x66, 0x33))
    return para


def add_table_custom(doc, headers, rows, col_widths=None):
    """Add styled table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, '4472C4')
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, bold=True, color=(0xFF, 0xFF, 0xFF))

    # Data rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = str(cell_text)
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    set_run_font(run)

    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(width)

    return table


def main():
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    # ====== Title Page ======
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_before = Pt(120)
    title_run = title_para.add_run('需求文档驱动测试用例生成系统')
    set_run_font(title_run, size=26, bold=True, color=(0x44, 0x72, 0xC4))

    sub_title = doc.add_paragraph()
    sub_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_title.paragraph_format.space_before = Pt(20)
    sub_run = sub_title.add_run('操作使用手册')
    set_run_font(sub_run, size=18, bold=True, color=(0x66, 0x66, 0x66))

    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.paragraph_format.space_before = Pt(60)
    vrun = version.add_run('Version 1.0\n2026年5月')
    set_run_font(vrun, size=12, color=(0x99, 0x99, 0x99))

    doc.add_page_break()

    # ====== Table of Contents ======
    add_heading_custom(doc, '目录', level=1)
    toc_items = [
        ('一、系统概述', 1),
        ('二、系统架构', 1),
        ('三、环境安装与配置', 1),
        ('四、核心功能使用指南', 1),
        ('  4.1 需求文档解析', 2),
        ('  4.2 需求模块拆分', 2),
        ('  4.3 测试用例生成', 2),
        ('  4.4 Excel 导出', 2),
        ('  4.5 文档变更检测', 2),
        ('  4.6 用例自动同步', 2),
        ('  4.7 规则校验', 2),
        ('  4.8 质量评分', 2),
        ('五、MCP Server 注册与使用', 1),
        ('六、数据目录结构', 1),
        ('七、常见问题', 1),
        ('八、附录', 1),
    ]
    for text, level in toc_items:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.5 * level)
        run = para.add_run(text)
        set_run_font(run, size=11 if level == 1 else 10)

    doc.add_page_break()

    # ====== Chapter 1 ======
    add_heading_custom(doc, '一、系统概述', level=1)
    add_paragraph_custom(doc,
        '本系统是一款基于 AI 大模型的测试用例自动生成工具，专为测试工程师设计。'
        '系统支持解析 PDF、Excel、Word、图片等多种格式的需求文档，运用资深测试专家的系统化方法论，'
        '自动生成高覆盖率、高质量的功能测试用例。')

    add_paragraph_custom(doc,
        '系统核心能力包括：多格式文档解析、需求模块拆分、AI 智能用例生成、'
        'Excel 标准化导出、需求变更检测、用例自动同步、规则校验和质量评分。')

    add_heading_custom(doc, '1.1 设计原则', level=2)
    principles = [
        '仅聚焦功能测试范畴，不生成性能、安全相关测试内容',
        '运用七大系统化测试设计方法（等价类、边界值、决策表、状态转换、正交试验、场景法、错误推测）',
        '高覆盖率保障：需求覆盖率100%、边界覆盖率100%、业务规则覆盖率100%、功能覆盖率≥95%、异常覆盖率≥90%',
        '模块化设计思路，保证结构清晰，支持长期维护优化',
        '预留测试空间，覆盖突发情况与非常规测试场景',
    ]
    for p in principles:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(p)
        set_run_font(run)

    # ====== Chapter 2 ======
    add_heading_custom(doc, '二、系统架构', level=1)
    add_paragraph_custom(doc, '系统采用模块化分层架构，包含以下核心组件：')

    add_table_custom(doc,
        ['组件', '文件路径', '职责'],
        [
            ['文档解析器', 'src/parsers/', 'PDF/Excel/Word/图片 → Markdown'],
            ['需求拆分器', 'src/analyzers/', '大需求 → 模块化子需求'],
            ['用例生成引擎', 'src/generators/', 'LLM 生成结构化测试用例 JSON'],
            ['Excel 导出器', 'src/exporters/', '用例 → 标准化 Excel'],
            ['差异检测引擎', 'src/diff_engine/', '语义级文档 Diff + 影响分析'],
            ['规则校验引擎', 'src/validators/', '12 条内置规则自动校验'],
            ['质量评分器', 'src/validators/', '5 维度质量评分'],
            ['MCP Server', 'mcp_server/server.py', '10 Tools + 3 Resources + 3 Prompts'],
        ],
        [3.5, 4.5, 7]
    )

    doc.add_paragraph()
    add_paragraph_custom(doc, '数据流转路径：')
    add_code_block(doc,
        '需求文档(PDF/Excel/Word/图片)\n'
        '        ↓\n'
        '  parsers/*.py  → 统一 Markdown\n'
        '        ↓\n'
        '  analyzers/requirement_splitter.py → 按模块拆分子需求\n'
        '        ↓\n'
        '  generators/case_generator.py → LLM 生成用例 JSON\n'
        '        ↓\n'
        '  exporters/excel_exporter.py → 标准化 Excel\n'
        '        ↓\n'
        '  validators/*.py → 规则校验 + 质量评分')

    # ====== Chapter 3 ======
    add_heading_custom(doc, '三、环境安装与配置', level=1)

    add_heading_custom(doc, '3.1 依赖安装', level=2)
    add_paragraph_custom(doc, '确保 Python 3.10+ 环境，执行以下命令安装依赖：')
    add_code_block(doc, 'pip install fastmcp anthropic pymupdf openpyxl pandas python-docx pillow pydantic httpx')

    add_heading_custom(doc, '3.2 API Key 配置', level=2)
    add_paragraph_custom(doc,
        '系统依赖 Claude API 进行需求分析和用例生成。请设置 Anthropic API Key：')
    add_code_block(doc, 'set ANTHROPIC_API_KEY=your_api_key_here')
    add_paragraph_custom(doc,
        '或在调用 generate_test_cases 时传入 api_key 参数。')

    add_heading_custom(doc, '3.3 MCP Server 注册', level=2)
    add_paragraph_custom(doc, '将 MCP Server 注册到 Claude Code 中，可通过自然语言调用所有功能：')
    add_code_block(doc, 'claude mcp add test-case-generator uv --directory e:/test_ai mcp_server.server:main')

    # ====== Chapter 4 ======
    add_heading_custom(doc, '四、核心功能使用指南', level=1)

    # 4.1
    add_heading_custom(doc, '4.1 需求文档解析', level=2)
    add_paragraph_custom(doc,
        '将需求文档（PDF、Excel、Word、图片）解析为结构化的 Markdown 文本，保留标题层级、表格、图片说明。')
    add_paragraph_custom(doc, '调用方式：', bold=True)
    add_code_block(doc, 'parse_document(file_path="path/to/requirement.pdf")')
    add_paragraph_custom(doc, '支持格式：')
    formats = [
        'PDF (.pdf) — 提取文本、表格、图片位置，按字体大小识别标题层级',
        'Excel (.xlsx, .xls) — 读取多 Sheet、合并单元格',
        'Word (.docx) — 保留标题层级（Heading 1/2/3），转换表格',
        '图片 (.png, .jpg, .jpeg) — Claude Vision API 高精度 OCR',
    ]
    for f in formats:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(f)
        set_run_font(run)
    add_paragraph_custom(doc, '输出：解析后的 Markdown 保存到 data/parsed/{doc_name}.md')

    # 4.2
    add_heading_custom(doc, '4.2 需求模块拆分', level=2)
    add_paragraph_custom(doc,
        '将大需求文档按标题层级拆分为多个独立的子需求模块，每个模块生成独立的 .md 文件和索引。')
    add_paragraph_custom(doc, '调用方式：', bold=True)
    add_code_block(doc, 'split_requirement(doc_name="requirement")')
    add_paragraph_custom(doc, '拆分策略：')
    strategies = [
        '一级/二级标题作为模块边界',
        '三级及以下作为子需求',
        '保留完整上下文和溯源信息（ID、模块名、优先级、依赖关系）',
    ]
    for s in strategies:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(s)
        set_run_font(run)
    add_paragraph_custom(doc, '输出：data/parsed/split/{doc_name}/ 目录下的独立 Markdown 文件 + 索引 JSON')

    # 4.3
    add_heading_custom(doc, '4.3 测试用例生成', level=2)
    add_paragraph_custom(doc,
        '基于解析后的需求文档，调用 Claude API 生成结构化测试用例 JSON。支持历史用例注入以保持风格一致性。')
    add_paragraph_custom(doc, '调用方式：', bold=True)
    add_code_block(doc,
        '# 整体生成\ngenerate_test_cases(doc_name="requirement")\n\n'
        '# 基于拆分后的模块逐个生成（推荐大型文档）\n'
        'generate_test_cases(doc_name="requirement", use_split=True)')
    add_paragraph_custom(doc, '生成策略：')
    gen_strategies = [
        '运用七大测试设计方法：等价类、边界值、决策表、状态转换、正交试验、场景法、错误推测',
        '高覆盖率保障：需求100%、边界100%、规则100%、功能≥95%、异常≥90%',
        '测试数据覆盖：正常合规、边界临界、异常违规、特殊格式',
        '操作行为覆盖：常规顺序、重复操作、逆向操作、随机顺序',
        '仅生成功能测试用例，不生成性能/安全用例',
        '每条用例标注使用的测试设计方法（test_method 字段）',
    ]
    for s in gen_strategies:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(s)
        set_run_font(run)
    add_paragraph_custom(doc, '输出：data/test_cases/{doc_name}_cases.json')

    # 4.4
    add_heading_custom(doc, '4.4 Excel 导出', level=2)
    add_paragraph_custom(doc,
        '将生成的测试用例导出为标准化 Excel 文件，支持多 Sheet 导出和优先级颜色标记。')
    add_paragraph_custom(doc, '调用方式：', bold=True)
    add_code_block(doc,
        '# 导出到默认路径\nexport_to_excel(doc_name="requirement")\n\n'
        '# 指定输出路径\n'
        'export_to_excel(doc_name="requirement", output_path="output.xlsx")')
    add_paragraph_custom(doc, 'Excel 模板包含列：用例ID、所属模块、用例标题、前置条件、测试步骤、预期结果、优先级、用例类型、需求来源、状态、文档版本')

    # 4.5
    add_heading_custom(doc, '4.5 文档变更检测', level=2)
    add_paragraph_custom(doc,
        '语义级检测需求文档的变更，对比新旧版本，识别新增/删除/修改的章节，输出结构化差异报告。')
    add_paragraph_custom(doc, '调用方式：', bold=True)
    add_code_block(doc, 'detect_doc_changes(doc_name="requirement", new_file_path="path/to/v2.pdf")')
    add_paragraph_custom(doc, '检测维度：')
    detects = [
        '章节新增/删除：按 #/##/### 标题拆分，ID 匹配',
        '内容修改：3-gram Jaccard 相似度，阈值 0.7',
        '表格变更：统计 | 分隔行数变化',
        '图片变更：统计 ![...] 和 [Image N] 标记数变化',
        '风险评级：有删除=高，大量增改=中，少量修改=低',
    ]
    for d in detects:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(d)
        set_run_font(run)
    add_paragraph_custom(doc, '输出：变更报告 Markdown + 差异数据 JSON')

    # 4.6
    add_heading_custom(doc, '4.6 用例自动同步', level=2)
    add_paragraph_custom(doc,
        '根据文档变更检测结果，自动分析对现有测试用例的影响，标记用例状态并生成同步报告。')
    add_paragraph_custom(doc, '调用方式：', bold=True)
    add_code_block(doc, 'sync_test_cases(doc_name="requirement", new_doc_name="requirement_v2")')
    add_paragraph_custom(doc, '同步策略：')
    sync = [
        '新增章节 → 标记需新增用例',
        '删除章节 → 关联用例标记 deprecated',
        '大幅修改(<50% 相似) → 关联用例标记 needs-review (high)',
        '轻微修改 → 关联用例标记 needs-review (normal)',
    ]
    for s in sync:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(s)
        set_run_font(run)
    add_paragraph_custom(doc, '输出：同步后用例 JSON + 同步报告 Markdown')

    # 4.7
    add_heading_custom(doc, '4.7 规则校验', level=2)
    add_paragraph_custom(doc,
        '对测试用例进行自动化规则校验，检查必填字段、格式规范、内容模式等，输出违规详情和修复建议。')
    add_paragraph_custom(doc, '调用方式：', bold=True)
    add_code_block(doc, 'validate_cases(doc_name="requirement")')
    add_paragraph_custom(doc, '内置规则（12 条）：')

    add_table_custom(doc,
        ['规则ID', '名称', '级别', '说明'],
        [
            ['R001', '用例ID必填', 'Error', '每条用例必须有唯一ID'],
            ['R002', '用例标题必填', 'Error', '标题不少于5个字符'],
            ['R003', '测试步骤必填', 'Error', '必须有具体步骤'],
            ['R004', '预期结果必填', 'Error', '必须有明确预期结果'],
            ['R005', '用例标题格式', 'Warning', '建议以动词开头'],
            ['R006', '步骤应具体可执行', 'Warning', '避免笼统描述'],
            ['R007', '预期结果应可验证', 'Warning', '需包含可观察关键词'],
            ['R008', '优先级格式正确', 'Error', '必须是 P0/P1/P2/P3'],
            ['R009', '用例类型有效', 'Warning', '使用标准类型名称'],
            ['R010', '应有需求来源', 'Warning', '记录来源文档章节'],
            ['R011', '前置条件建议填写', 'Info', '复杂用例建议填写'],
            ['R012', '应有模块归属', 'Warning', '指定所属模块'],
        ],
        [2, 4, 2.5, 6.5]
    )

    # 4.8
    add_heading_custom(doc, '4.8 质量评分', level=2)
    add_paragraph_custom(doc,
        '从五个维度对测试用例进行质量评分，输出总体评分、维度得分和改进建议。')
    add_paragraph_custom(doc, '调用方式：', bold=True)
    add_code_block(doc, 'score_case_quality(doc_name="requirement")')
    add_paragraph_custom(doc, '评分维度：')

    add_table_custom(doc,
        ['维度', '权重', '评分逻辑'],
        [
            ['覆盖率', '25%', '类型多样性、异常场景、边界值、test_method 字段'],
            ['可执行性', '25%', '步骤数量、操作动词、可验证关键词、结果长度'],
            ['准确性', '20%', '无歧义词汇、步骤逻辑矛盾检测'],
            ['完整性', '20%', '前置条件率、模块归属率、需求溯源率'],
            ['优先级', '10%', 'P0 占比合理性、梯度分布'],
        ],
        [3, 2.5, 9.5]
    )

    add_paragraph_custom(doc, '等级评定：A+ (4.5+) / A (4.0+) / B+ (3.5+) / B (3.0+) / C (2.5+) / D (2.0+) / F (<2.0)')

    # ====== Chapter 5 ======
    add_heading_custom(doc, '五、MCP Server 注册与使用', level=1)
    add_paragraph_custom(doc,
        '注册后可在 Claude Code 中通过自然语言调用所有功能，无需记忆命令。')

    add_heading_custom(doc, '5.1 注册命令', level=2)
    add_code_block(doc, 'claude mcp add test-case-generator uv --directory e:/test_ai mcp_server.server:main')

    add_heading_custom(doc, '5.2 自然语言调用示例', level=2)
    examples = [
        '帮我解析这个 PDF 需求文档并生成测试用例',
        '将需求文档按模块拆分，然后逐个生成用例',
        '把生成的测试用例导出成 Excel',
        '检测需求文档是否有变更，更新对应的测试用例',
        '校验刚才生成的测试用例质量',
        '给这批用例打个质量分',
    ]
    for ex in examples:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(ex)
        set_run_font(run)

    add_heading_custom(doc, '5.3 完整 API 列表', level=2)
    add_table_custom(doc,
        ['类型', '名称', '功能描述'],
        [
            ['Tool', 'parse_document', '解析 PDF/Excel/Word/图片 → Markdown'],
            ['Tool', 'split_requirement', '按模块拆分子需求'],
            ['Tool', 'generate_test_cases', 'LLM 生成测试用例 JSON'],
            ['Tool', 'export_to_excel', '导出标准化 Excel'],
            ['Tool', 'detect_doc_changes', '语义级变更检测'],
            ['Tool', 'sync_test_cases', '自动同步用例状态'],
            ['Tool', 'validate_cases', '规则校验'],
            ['Tool', 'score_case_quality', '质量评分'],
            ['Tool', 'list_documents', '列出已解析文档'],
            ['Tool', 'list_test_cases', '列出已生成用例'],
            ['Resource', 'test-case-library://', '完整用例库 JSON'],
            ['Resource', 'requirement-doc://{name}', '需求文档 Markdown'],
            ['Resource', 'requirement-modules://{name}', '子需求索引'],
            ['Prompt', 'analyze-requirement', '需求分析'],
            ['Prompt', 'generate-boundary-cases', '边界值生成'],
            ['Prompt', 'review-case-quality', '质量评审'],
        ],
        [2.5, 4, 8.5]
    )

    # ====== Chapter 6 ======
    add_heading_custom(doc, '六、数据目录结构', level=1)
    add_paragraph_custom(doc, '系统数据按以下目录结构组织：')
    add_code_block(doc,
        'data/\n'
        '├── requirements/          # 原始需求文档（用户上传）\n'
        '├── parsed/                # 解析后的结构化文本\n'
        '│   ├── {doc_name}.md      # 完整文档 Markdown\n'
        '│   └── split/             # 拆分后的子需求\n'
        '│       └── {doc_name}/\n'
        '│           ├── {doc_name}_index.json\n'
        '│           └── {doc_name}_*.md\n'
        '└── test_cases/            # 生成的测试用例\n'
        '    ├── {doc_name}_cases.json\n'
        '    ├── {doc_name}_test_cases.xlsx\n'
        '    ├── {doc_name}_diff.json\n'
        '    ├── {doc_name}_change_report.md\n'
        '    ├── {doc_name}_impact.json\n'
        '    ├── {doc_name}_sync_report.md\n'
        '    ├── {doc_name}_validation.json\n'
        '    └── {doc_name}_quality.json')

    # ====== Chapter 7 ======
    add_heading_custom(doc, '七、常见问题', level=1)

    faqs = [
        ('Q: 解析图片需求文档时提示 API Key 错误？',
         'A: 图片解析依赖 Claude Vision API，需要设置 ANTHROPIC_API_KEY 环境变量。'),
        ('Q: 生成的用例质量不高怎么办？',
         'A: 1) 确保需求文档解析完整；2) 提供历史用例 Excel 到 data/historical_cases/ 目录作为风格参考；3) 使用 score_case_quality 查看评分详情并针对性优化。'),
        ('Q: 大型需求文档如何处理？',
         'A: 建议先用 split_requirement 拆分文档，再用 generate_test_cases 的 use_split=True 参数逐个模块生成。'),
        ('Q: 需求变更后如何更新用例？',
         'A: 1) detect_doc_changes 检测变更；2) sync_test_cases 分析影响并标记状态；3) 根据报告中的建议操作（新增/Review/废弃）。'),
        ('Q: 如何扩展自定义校验规则？',
         'A: 在 RuleEngine 的 _build_builtin_rules 方法中添加新的 Rule 对象，或调用 add_rule() 动态添加。'),
    ]
    for q, a in faqs:
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(12)
        qrun = para.add_run(q)
        set_run_font(qrun, bold=True)
        para2 = doc.add_paragraph()
        para2.paragraph_format.left_indent = Cm(0.5)
        arun = para2.add_run(a)
        set_run_font(arun)

    # ====== Chapter 8 ======
    add_heading_custom(doc, '八、附录', level=1)

    add_heading_custom(doc, '8.1 七大测试设计方法速查', level=2)
    add_table_custom(doc,
        ['方法', '适用场景', '核心要点'],
        [
            ['等价类划分', '输入验证类功能', '有效等价类 + 无效等价类'],
            ['边界值分析', '数值/长度/范围限制', '最小值、最小值+1、正常值、最大值-1、最大值'],
            ['决策表', '多条件组合', '条件桩 + 动作桩，覆盖所有规则组合'],
            ['状态转换', '有状态流转的功能', '初始/中间/结束/异常状态及转换路径'],
            ['正交试验', '多因素交叉', '提取关键因子，精简组合数量'],
            ['场景法', '业务流程', '主流程 + 异常分支 + 边界极限'],
            ['错误推测', '高风险模块', '历史缺陷 + 业务风险 + 用户误操作'],
        ],
        [3, 4.5, 7.5]
    )

    add_heading_custom(doc, '8.2 优先级定义', level=2)
    add_table_custom(doc,
        ['优先级', '定义', '覆盖要求'],
        [
            ['P0', '核心业务流程、阻塞性缺陷、高风险区域', '必须100%覆盖'],
            ['P1', '重要功能节点、常见用户场景', '必须覆盖'],
            ['P2', '辅助功能、一般操作场景', '选择性覆盖'],
            ['P3', '边缘场景、低频操作、异常恢复', '抽样覆盖'],
        ],
        [2.5, 6.5, 6]
    )

    add_heading_custom(doc, '8.3 联系方式', level=2)
    add_paragraph_custom(doc,
        '如有问题或建议，请通过以下方式反馈：\n'
        '• GitHub Issues: https://github.com/your-org/test-case-generator/issues\n'
        '• 邮箱: test-team@example.com')

    # Save
    output_path = 'e:/test_ai/操作使用手册.docx'
    doc.save(output_path)
    print(f'文档已生成: {output_path}')


if __name__ == '__main__':
    main()
