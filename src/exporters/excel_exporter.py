"""测试用例 Excel 导出器."""

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """将测试用例导出为标准化 Excel 文件."""

    # 列定义
    COLUMNS = [
        ("用例ID", "id", 15),
        ("所属模块", "module", 20),
        ("用例标题", "title", 40),
        ("前置条件", "precondition", 30),
        ("测试步骤", "steps", 50),
        ("预期结果", "expected", 40),
        ("优先级", "priority", 10),
        ("用例类型", "type", 12),
        ("需求来源", "source", 25),
        ("状态", "status", 10),
        ("文档版本", "doc_version", 12),
    ]

    def __init__(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "测试用例"

    def export(
        self,
        cases: list[dict[str, Any]],
        output_path: str | Path,
        sheet_name: str | None = None,
    ) -> Path:
        """导出测试用例到 Excel.

        Args:
            cases: 测试用例列表
            output_path: 输出文件路径
            sheet_name: 工作表名称（可选）

        Returns:
            输出文件路径
        """
        if sheet_name:
            self.ws.title = sheet_name

        self._write_header()
        self._write_data(cases)
        self._apply_styles()

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(str(output_path))
        return output_path

    def export_multi_sheets(
        self,
        module_cases: dict[str, list[dict[str, Any]]],
        output_path: str | Path,
    ) -> Path:
        """按模块导出到多个工作表.

        Args:
            module_cases: {模块名: [用例列表]}
            output_path: 输出文件路径

        Returns:
            输出文件路径
        """
        first = True
        for module_name, cases in module_cases.items():
            if first:
                ws = self.ws
                ws.title = module_name[:31]  # Excel sheet name max 31 chars
                first = False
            else:
                ws = self.wb.create_sheet(title=module_name[:31])

            self.ws = ws
            self._write_header()
            self._write_data(cases)
            self._apply_styles()

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(str(output_path))
        return output_path

    def _write_header(self) -> None:
        """写入表头."""
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, (title, _, width) in enumerate(self.COLUMNS, start=1):
            cell = self.ws.cell(row=1, column=col_idx, value=title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            self.ws.column_dimensions[get_column_letter(col_idx)].width = width

        self.ws.row_dimensions[1].height = 30

    def _write_data(self, cases: list[dict[str, Any]]) -> None:
        """写入用例数据."""
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for row_idx, case in enumerate(cases, start=2):
            for col_idx, (_, field, _) in enumerate(self.COLUMNS, start=1):
                value = case.get(field, "")

                # 处理列表类型（如 steps）
                if isinstance(value, list):
                    value = "\n".join(f"{i+1}. {step}" for i, step in enumerate(value))

                cell = self.ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

                # 优先级颜色标记
                if field == "priority":
                    cell.font = self._priority_font(value)

            # 根据内容自动调整行高
            self.ws.row_dimensions[row_idx].height = max(30, len(str(case.get("steps", ""))) // 2)

    def _apply_styles(self) -> None:
        """应用全局样式."""
        # 冻结首行
        self.ws.freeze_panes = "A2"

        # 自动筛选
        self.ws.auto_filter.ref = self.ws.dimensions

    def _priority_font(self, priority: str) -> Font:
        """根据优先级返回字体样式."""
        colors = {
            "P0": "C00000",  # 红色
            "P1": "ED7D31",  # 橙色
            "P2": "FFC000",  # 黄色
            "P3": "70AD47",  # 绿色
        }
        color = colors.get(str(priority).upper(), "000000")
        return Font(bold=True, color=color)

    @staticmethod
    def read_historical_cases(path: str | Path) -> list[dict[str, Any]]:
        """读取历史用例 Excel 文件.

        Args:
            path: Excel 文件路径

        Returns:
            用例字典列表
        """
        from openpyxl import load_workbook

        wb = load_workbook(str(path))
        ws = wb.active

        # 读取表头
        headers = [cell.value for cell in ws[1]]
        field_map = {
            "用例ID": "id",
            "所属模块": "module",
            "用例标题": "title",
            "前置条件": "precondition",
            "测试步骤": "steps",
            "预期结果": "expected",
            "优先级": "priority",
            "用例类型": "type",
            "需求来源": "source",
        }

        cases: list[dict[str, Any]] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            case: dict[str, Any] = {}
            for header, value in zip(headers, row):
                if header in field_map and value is not None:
                    field = field_map[header]
                    # 步骤可能是多行文本
                    if field == "steps" and isinstance(value, str):
                        value = [s.strip() for s in value.split("\n") if s.strip()]
                    case[field] = value
            if case:
                cases.append(case)

        wb.close()
        return cases
